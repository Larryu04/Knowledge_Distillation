"""
Stage 1 — Syllabus ingestion.

Uploaded files (PDF / DOCX / TXT / MD) and pasted text become plain text, then the teacher
LLM turns each source into structured courses: a name, a short English domain label, the
granular topic list the syllabus actually covers, and terms that are explicitly out of scope.

That topic list is the raw material for generation: every synthetic question is anchored to
one topic, which is what keeps a 1,000-row course on-syllabus instead of drifting.
"""

import json
import os
import re

from .. import db
from ..jobs import estimate
from ..llm.teacher import make_teacher
from ..outline import parse_outline, strip_unit_prefix
from .base import StageContext, StageError, atomic_write_json, read_json, slugify

MAX_SOURCE_CHARS = 100_000      # ~30-50k tokens; well inside gpt-4o-mini's context
MIN_TOPICS = 5

SYSTEM_PROMPT = """COURSE_EXTRACTION
You convert university syllabus documents into a structured JSON description used to build \
a course-specific AI teaching assistant.

Rules:
1. A document may describe ONE course or SEVERAL courses (e.g. a whole degree program). \
Return one entry per distinct course.
2. "topics" must be granular, concrete subtopics taken ONLY from the syllabus (lecture \
topics, weekly units, learning outcomes). Split broad units into specific teachable points. \
Aim for 12-60 topics per course, but a short outline is still a course: a plain list such as \
"Week 1: X / Week 2: Y" gives one topic per line (keep its wording, drop the "Week N:" prefix). \
Never invent topics the syllabus does not cover.
3. "out_of_scope": terms the syllabus explicitly defers to other/later courses or excludes. \
Empty list if none are stated. Do not guess.
4. "style_note": only if the course needs special handling (e.g. a proof-heavy math course: \
"prefer verbal, conceptual explanations over long formal proofs"). Otherwise null.
5. Keep course names and topics in the syllabus's own language. "domain" is a short English \
label (2-5 words) for the subject.
6. If the text names no course title, set "course_name" to "" rather than inventing one.
7. If the text is not a syllabus at all, return {"courses": []}.

Return ONLY a JSON object:
{"courses": [{"course_name": str, "domain": str, "language": str, "summary": str, \
"topics": [str, ...], "out_of_scope": [str, ...], "style_note": str|null}]}"""


# --------------------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------------------
def extract_text(path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1255", "cp1256", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------------------
# Structuring
# --------------------------------------------------------------------------------------
def _str_list(val) -> list[str]:
    if isinstance(val, str):
        val = [ln for ln in re.split(r"[\n;]", val)]
    if not isinstance(val, list):
        return []
    out, seen = [], set()
    for item in val:
        s = strip_unit_prefix(re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", str(item)))
        if s and s.lower() not in seen:
            seen.add(s.lower())
            out.append(s)
    return out


def normalize_courses(raw_courses: list, source_name: str, fallback_name: str = "") -> list[dict]:
    """Untitled courses get `fallback_name` (numbered if a source holds several), else the
    source name, so an outline without a "Course:" line is still usable."""
    courses, untitled = [], 0
    for c in raw_courses or []:
        if not isinstance(c, dict):
            continue
        name = str(c.get("course_name") or "").strip()
        topics = _str_list(c.get("topics"))
        if len(topics) < MIN_TOPICS:
            continue
        if not name:
            untitled += 1
            name = fallback_name or source_name or "Course"
            if untitled > 1:
                name = f"{name} {untitled}"
        style = c.get("style_note")
        courses.append({
            "course_name": name,
            "domain": str(c.get("domain") or name).strip()[:60],
            "language": str(c.get("language") or "").strip(),
            "summary": str(c.get("summary") or "").strip(),
            "topics": topics,
            "out_of_scope": _str_list(c.get("out_of_scope")),
            "style_note": style.strip() if isinstance(style, str) and style.strip() else None,
            "source": source_name,
        })
    return courses


def dedupe_courses(courses: list[dict]) -> list[dict]:
    """Same course in two files -> merge topic lists; distinct courses with clashing domain
    labels -> suffix the label (domain is the per-course key for metrics and sampling)."""
    by_name: dict[str, dict] = {}
    for c in courses:
        key = c["course_name"].casefold()
        if key in by_name:
            existing = by_name[key]
            seen = {t.casefold() for t in existing["topics"]}
            existing["topics"] += [t for t in c["topics"] if t.casefold() not in seen]
            existing["out_of_scope"] = sorted(set(existing["out_of_scope"]) | set(c["out_of_scope"]))
        else:
            by_name[key] = dict(c)
    result, used_domains, used_slugs = [], set(), set()
    for c in by_name.values():
        domain, n = c["domain"], 2
        while domain.casefold() in used_domains:
            domain = f"{c['domain']} {n}"
            n += 1
        used_domains.add(domain.casefold())
        c["domain"] = domain
        slug = slugify(domain, fallback=f"course-{len(result) + 1}")
        base, n = slug, 2
        while slug in used_slugs:
            slug = f"{base}-{n}"
            n += 1
        used_slugs.add(slug)
        c["slug"] = slug
        result.append(c)
    return result


def run(ctx: StageContext) -> None:
    sources = read_json(ctx.ws.sources, default=[])
    if not sources:
        raise StageError("No syllabus files or pasted text were provided.")

    teacher = make_teacher(os.environ.get("GLOBAL_TEACHER_API_KEY"), ctx.config.teacher_model, log=ctx.log,
                           demo=ctx.config.demo)
    ctx.reporter.start(f"Reading {len(sources)} source(s)")

    all_courses = []
    for i, src in enumerate(sources):
        path = ctx.ws.uploads / src["file"]
        try:
            text = clean_text(extract_text(path))
        except Exception as e:  # noqa: BLE001 - a broken file should name itself, not crash opaquely
            raise StageError(f"Could not read {src['name']}: {type(e).__name__}: {e}") from e
        if len(text) < 200 and not parse_outline(text, MIN_TOPICS):
            ctx.log(f"Skipping {src['name']}: only {len(text)} characters of text "
                    "(scanned PDFs need OCR before upload).")
            ctx.reporter.update(i + 1, len(sources), f"Skipped {src['name']}")
            continue
        if len(text) > MAX_SOURCE_CHARS:
            ctx.log(f"{src['name']}: {len(text):,} characters; using the first {MAX_SOURCE_CHARS:,}.")
            text = text[:MAX_SOURCE_CHARS]

        ctx.reporter.update(i, len(sources), f"Structuring {src['name']}", force=True)
        user = f"# Source: {src['name']}\n\n{text}"
        # One source -> the program name is the natural title; several -> say which source it was.
        fallback = (ctx.config.program_name if len(sources) == 1
                    else f"{ctx.config.program_name} ({src['name']})")
        courses = []
        for attempt in range(3):
            raw = teacher.complete_json(SYSTEM_PROMPT, user, temperature=0.2)
            try:
                courses = normalize_courses(json.loads(raw).get("courses"), src["name"], fallback)
                break
            except (json.JSONDecodeError, AttributeError):
                ctx.log(f"{src['name']}: malformed extraction (attempt {attempt + 1}/3), retrying")
        if not courses and (outline := parse_outline(text, MIN_TOPICS)):
            # The teacher found nothing usable in a short or informal outline: read the list as-is.
            courses = normalize_courses(outline, src["name"], fallback)
            ctx.log(f"{src['name']}: read as a plain topic outline")
        if not courses:
            ctx.log(f"No course with at least {MIN_TOPICS} topics found in {src['name']}.")
        for c in courses:
            ctx.log(f"Found course '{c['course_name']}' ({len(c['topics'])} topics) in {src['name']}")
        all_courses += courses
        ctx.reporter.update(i + 1, len(sources), f"Parsed {src['name']}", force=True)

    courses = dedupe_courses(all_courses)
    if not courses:
        raise StageError("No usable courses were found in the provided syllabi. Each course needs a "
                         f"readable topic list (at least {MIN_TOPICS} topics).")

    atomic_write_json(ctx.ws.syllabi, {"program": ctx.config.program_name, "courses": courses})

    # Refresh the ETA / cost estimate now that the real course count is known.
    db.update_job(ctx.job_id, estimate=estimate(ctx.config, len(courses), len(sources)))
    ctx.reporter.finish(f"{len(courses)} course(s), {sum(len(c['topics']) for c in courses)} topics",
                        usage=teacher.usage.as_dict(), courses=len(courses))
